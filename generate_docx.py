import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page Margins: 1 inch on all sides
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    primary_color = RGBColor(79, 70, 229)   # Indigo #4F46E5
    dark_color = RGBColor(15, 23, 42)       # Slate #0F172A
    gray_color = RGBColor(100, 116, 139)    # Slate #64748B

    # -------------------------------------------------------------
    # COVER / TITLE PAGE
    # -------------------------------------------------------------
    p_title_space = doc.add_paragraph()
    p_title_space.paragraph_format.space_before = Pt(36)

    p_proj_title = doc.add_paragraph()
    p_proj_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_proj = p_proj_title.add_run("STUDENT ATTENDANCE MANAGEMENT\nMOBILE APPLICATION")
    run_proj.font.name = "Arial"
    run_proj.font.size = Pt(24)
    run_proj.font.bold = True
    run_proj.font.color.rgb = primary_color

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("A Project Report submitted in partial fulfillment of the requirements for the award of the degree of")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = gray_color
    p_sub.paragraph_format.space_before = Pt(12)
    p_sub.paragraph_format.space_after = Pt(8)

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_deg = p_deg.add_run("MASTER OF COMPUTER APPLICATIONS\n(MCA)")
    run_deg.font.name = "Arial"
    run_deg.font.size = Pt(16)
    run_deg.font.bold = True
    run_deg.font.color.rgb = dark_color

    p_tech = doc.add_paragraph()
    p_tech.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_tech = p_tech.add_run("Developed Using React Native & Expo Framework (Offline-First Architecture)")
    run_tech.font.name = "Arial"
    run_tech.font.size = Pt(11)
    run_tech.font.color.rgb = primary_color
    p_tech.paragraph_format.space_after = Pt(48)

    # Submission metadata box
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_meta.add_run("DEPARTMENT OF COMPUTER APPLICATIONS\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r2 = p_meta.add_run("Academic Year 2026 - 2027\n\nProject Repository: D:\\1.MCA-PROJECTS\\student-attendance-app")
    r2.font.name = "Arial"
    r2.font.size = Pt(10)
    r2.font.color.rgb = gray_color

    doc.add_page_break()

    # -------------------------------------------------------------
    # CERTIFICATE & DECLARATION
    # -------------------------------------------------------------
    p_cert_title = doc.add_paragraph()
    p_cert_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cert = p_cert_title.add_run("CERTIFICATE OF BONAFIDE WORK")
    run_cert.font.name = "Arial"
    run_cert.font.size = Pt(16)
    run_cert.font.bold = True
    run_cert.font.color.rgb = dark_color
    p_cert_title.paragraph_format.space_after = Pt(18)

    p_cert_body = doc.add_paragraph()
    p_cert_body.paragraph_format.line_spacing = 1.25
    r_cb = p_cert_body.add_run(
        "This is to certify that the project entitled \"STUDENT ATTENDANCE MANAGEMENT MOBILE APPLICATION\" "
        "is a bonafide work carried out by the student in partial fulfillment of the requirements for the award "
        "of the degree of Master of Computer Applications (MCA). It is confirmed that the application has been "
        "thoroughly designed, implemented, and tested in accordance with the prescribed academic standards."
    )
    r_cb.font.name = "Arial"
    r_cb.font.size = Pt(11)
    p_cert_body.paragraph_format.space_after = Pt(40)

    # Signatures table
    sig_table = doc.add_table(rows=2, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_headers = ["Internal Guide", "Head of Department", "External Examiner"]
    for i, title in enumerate(sig_headers):
        cell_top = sig_table.cell(0, i)
        cell_top.paragraphs[0].text = "\n\n____________________\n" + title
        cell_top.paragraphs[0].runs[0].font.name = "Arial"
        cell_top.paragraphs[0].runs[0].font.size = Pt(10)
        cell_top.paragraphs[0].runs[0].font.bold = True
        cell_top.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # -------------------------------------------------------------
    # ABSTRACT & ACKNOWLEDGEMENT
    # -------------------------------------------------------------
    p_abs_title = doc.add_paragraph()
    r_at = p_abs_title.add_run("ABSTRACT")
    r_at.font.name = "Arial"
    r_at.font.size = Pt(16)
    r_at.font.bold = True
    r_at.font.color.rgb = primary_color
    p_abs_title.paragraph_format.space_after = Pt(12)

    p_abs = doc.add_paragraph()
    p_abs.paragraph_format.line_spacing = 1.25
    r_ab = p_abs.add_run(
        "Manual attendance management across educational institutions has persistently resulted in substantial administrative overhead, "
        "frequent human tabulation errors, loss of lecture time, and delayed detection of students falling below mandatory university "
        "attendance requirements (the 75% minimum attendance rule for semester examinations).\n\n"
        "This project introduces the Student Attendance Management System, a production-grade cross-platform mobile application "
        "built with React Native and Expo. The system provides seamless student registration across multiple academic cohorts, "
        "intuitive single-tap roll call recording (Present, Absent, Late), dynamic percentage calculations, and persistent on-device "
        "storage utilizing AsyncStorage. A cornerstone feature of this project is the Examination Defaulter Recovery Engine, which "
        "mathematically derives the exact number of consecutive upcoming lectures an at-risk student must attend to regain examination "
        "eligibility. The solution operates completely offline without recurring cloud infrastructure costs, delivering high reliability, "
        "instantaneous query response times, and comprehensive analytical reporting."
    )
    r_ab.font.name = "Arial"
    r_ab.font.size = Pt(11)

    doc.add_page_break()

    # -------------------------------------------------------------
    # CHAPTERS
    # -------------------------------------------------------------
    chapters = [
        {
            "num": "1",
            "title": "INTRODUCTION & PROBLEM STATEMENT",
            "sections": [
                ("1.1 Project Overview",
                 "The Student Attendance Management Application is an administrative mobile solution tailored for collegiate faculty. "
                 "Engineered using React Native and the Expo SDK, the application converts the conventional paper-based roll-call register "
                 "into a highly responsive digital workspace."),
                ("1.2 Problem Statement",
                 "Traditional manual attendance poses significant limitations:\n"
                 "• Time Consumption: Calling out roll numbers daily consumes 10-15% of classroom teaching time.\n"
                 "• Arithmetic Errors: Aggregating attendance records at the end of semester often leads to miscalculated percentages.\n"
                 "• Delayed Defaulter Detection: Students lacking attendance are often identified only days before exams, eliminating recovery opportunities.\n"
                 "• Physical Degradation: Paper registers suffer from physical wear, tear, ink blurs, and risk of misplacement."),
                ("1.3 Objectives of the Project",
                 "• Develop an intuitive, mobile roll-call interface allowing 1-tap toggling between Present, Absent, and Late.\n"
                 "• Implement real-time mathematical aggregation of attendance percentages per student and cohort.\n"
                 "• Provide an automated Defaulter Tracking Engine highlighting students under the 75% attendance compliance rule.\n"
                 "• Formulate a Defaulter Recovery algorithm computing required consecutive attendances.\n"
                 "• Ensure robust local persistence via AsyncStorage with zero cloud dependencies.\n"
                 "• Deliver date-wise, student-wise, and summary reporting with native mobile sharing."),
                ("1.4 Scope & Target Audience",
                 "The application caters to university professors, department heads, and academic mentors. It is scalable across diverse "
                 "academic programs including MCA, BCA, B.Tech, and M.Tech.")
            ]
        },
        {
            "num": "2",
            "title": "LITERATURE SURVEY & SYSTEM COMPARISON",
            "sections": [
                ("2.1 Existing Systems vs Proposed System",
                 "A comparative assessment between traditional attendance modes and the proposed React Native application is summarized below:")
            ],
            "table": {
                "headers": ["Feature / Metric", "Paper Register", "Desktop Web Portal", "Proposed Mobile App"],
                "rows": [
                    ["Portability", "Bulky physical book", "Confined to PC desk", "High (Pocket smartphone)"],
                    ["Roll Call Speed", "Slow (pen marking)", "Moderate (keyboard)", "Instantaneous (1-tap toggle)"],
                    ["Offline Capability", "Yes", "No (Internet required)", "Yes (100% Offline AsyncStorage)"],
                    ["Calculation Speed", "Manual calculator", "Server batch job", "Real-Time instant math"],
                    ["Defaulter Detection", "End of term only", "Periodic reports", "Instant live notification"],
                    ["Recovery Guidance", "None", "None", "Dynamic consecutive formula"],
                    ["Infrastructure Cost", "Paper supplies", "Server & domain fees", "Zero hosting cost"]
                ]
            }
        },
        {
            "num": "3",
            "title": "SOFTWARE REQUIREMENT SPECIFICATIONS (SRS)",
            "sections": [
                ("3.1 Functional Requirements",
                 "• FR1 - Student Directory: Registration with unique Roll Number validation, Name, Class/Batch, Gender, Email, and Phone.\n"
                 "• FR2 - Daily Roll Call: Interactive marking screen with date selector, batch dropdown, and Present/Absent/Late status toggles.\n"
                 "• FR3 - Bulk Operations: Quick buttons for 'Mark All Present', 'Mark All Absent', and 'Clear'.\n"
                 "• FR4 - Analytical Engine: Automated calculation of aggregate percentage, attended count, and absent counts.\n"
                 "• FR5 - Defaulter Management: Automatic flagging of students with attendance < 75% and calculation of recovery targets.\n"
                 "• FR6 - Local Storage: Complete persistence of all datasets via AsyncStorage.\n"
                 "• FR7 - Reporting & Native Sharing: Exporting date-wise and cumulative records through the device Share sheet."),
                ("3.2 Non-Functional Requirements",
                 "• Performance: Screen rendering and calculations execute in under 100 milliseconds.\n"
                 "• Usability: High-contrast color standards (Emerald for Present, Crimson for Absent, Amber for Late).\n"
                 "• Reliability: Zero data loss during unexpected app termination via asynchronous transactional writes.\n"
                 "• Cross-Platform: 100% code sharing across Android, iOS, and Web.")
            ]
        },
        {
            "num": "4",
            "title": "SYSTEM ARCHITECTURE & DESIGN",
            "sections": [
                ("4.1 Layered Architecture Overview",
                 "The system follows a modern decoupled 4-tier mobile architecture:\n"
                 "1. Presentation Layer: React Native functional components, responsive UI cards, and custom badges.\n"
                 "2. Navigation Layer: React Navigation 7 featuring Bottom Tab Navigator and Native Stack Navigator.\n"
                 "3. State Management Layer: React Context API (AttendanceContext) providing reactive, centralized state.\n"
                 "4. Data Persistence Layer: StorageService abstraction communicating with AsyncStorage key-value engine."),
                ("4.2 Entity Relationship (ER) Data Model",
                 "• Student Entity: { id, rollNo, name, batch, email, phone, gender, createdAt }\n"
                 "• Attendance Record Entity: { sessionKey (Date_Batch), date, batch, subject, records: { [studentId]: status }, updatedAt }\n"
                 "• Batch Entity: Collection of academic cohort names (e.g., 'MCA 1st Year', 'MCA 2nd Year').")
            ]
        },
        {
            "num": "5",
            "title": "MATHEMATICAL FORMULATIONS & CORE ALGORITHMS",
            "sections": [
                ("5.1 Attendance Percentage Formulation",
                 "Let P denote total Present sessions, L denote Late sessions (weighted at 0.5), and A denote Absent sessions.\n"
                 "Total Sessions Conducted: N = P + L + A\n"
                 "Effective Sessions Attended: E = P + (0.5 * L)\n"
                 "Attendance Percentage (%) = (E / N) * 100\n\n"
                 "Compliance Status Categories:\n"
                 "• Percentage >= 75%: Safe (Eligible for Examinations)\n"
                 "• 65% <= Percentage < 75%: Warning (At Risk)\n"
                 "• Percentage < 65%: Critical (Examination Defaulter)"),
                ("5.2 Defaulter Recovery Formulation",
                 "To determine the minimum number of consecutive upcoming lectures (X) a student must attend without missing any to achieve at least 75% (0.75):\n"
                 "(E + X) / (N + X) >= 0.75\n"
                 "E + X >= 0.75 * (N + X)\n"
                 "0.25 * X >= 0.75 * N - E\n"
                 "X = ceil( (0.75 * N - E) / 0.25 )\n\n"
                 "This mathematical formulation provides faculty with an actionable metric when advising students during attendance counseling.")
            ]
        },
        {
            "num": "6",
            "title": "SOFTWARE TESTING & TEST CASES",
            "sections": [
                ("6.1 Testing Methodology",
                 "The application underwent comprehensive functional and user acceptance testing across both Android physical devices and web emulators.")
            ],
            "table": {
                "headers": ["Test ID", "Module", "Scenario", "Input", "Expected Output", "Status"],
                "rows": [
                    ["TC01", "Registration", "Empty student name", "Name='', Roll='MCA-99'", "Validation error alert displayed", "PASS"],
                    ["TC02", "Registration", "Duplicate roll number", "Roll='MCA-01'", "Duplicate roll error alert displayed", "PASS"],
                    ["TC03", "Roll Call", "Mark All Present", "Click 'All Present'", "All students toggled to Present", "PASS"],
                    ["TC04", "Roll Call", "Unmarked students prompt", "5 of 10 marked", "Alert offering auto-mark absent", "PASS"],
                    ["TC05", "Analytics", "Percentage with Late", "6 Present, 2 Late, 2 Absent", "E=7, N=10 -> 70.0% calculated", "PASS"],
                    ["TC06", "Defaulter", "Recovery Calculation", "E=7, N=10", "X = ceil((7.5-7)/0.25) = 2 classes", "PASS"],
                    ["TC07", "Persistence", "Restart retention", "Kill & restart app", "All students & sessions retained", "PASS"],
                    ["TC08", "Export", "Native Share trigger", "Click Share in Reports", "Device Share sheet opened with CSV text", "PASS"]
                ]
            }
        },
        {
            "num": "7",
            "title": "CONCLUSION & FUTURE ENHANCEMENTS",
            "sections": [
                ("7.1 Conclusion",
                 "The Student Attendance Management Mobile Application accomplishes all specified objectives. It successfully replaces manual "
                 "paper-based registers with an elegant, responsive, and robust mobile workflow. Through dynamic percentage computation, "
                 "examination defaulter tracking, and recovery calculations, the application enhances administrative efficiency in higher education."),
                ("7.2 Future Enhancements",
                 "• Biometric Facial Recognition: Automating roll call using device camera and face recognition APIs.\n"
                 "• Cloud Sync & Multi-Faculty Collaboration: Integrating Firebase or Supabase for real-time synchronization across devices.\n"
                 "• Automated Parent SMS / WhatsApp Notifications: Instant notification dispatch when a student is recorded absent.\n"
                 "• Timed QR Code Attendance: Displaying a dynamic rotating QR code on lecture slides for student self-check-in.")
            ]
        },
        {
            "num": "8",
            "title": "REFERENCES",
            "sections": [
                ("8.1 Bibliography",
                 "1. Facebook Open Source, 'React Native Documentation', https://reactnative.dev\n"
                 "2. Expo Documentation & Guides, https://docs.expo.dev\n"
                 "3. React Navigation v7 Core Architecture, https://reactnavigation.org\n"
                 "4. AsyncStorage Community Module, https://react-native-async-storage.github.io/async-storage\n"
                 "5. Pressman, R. S., 'Software Engineering: A Practitioner's Approach', McGraw-Hill Education, 8th Edition.")
            ]
        }
    ]

    for chap in chapters:
        # Chapter Heading
        p_ch = doc.add_paragraph()
        p_ch.paragraph_format.space_before = Pt(16)
        p_ch.paragraph_format.space_after = Pt(8)
        r_ch = p_ch.add_run(f"CHAPTER {chap['num']}: {chap['title']}")
        r_ch.font.name = "Arial"
        r_ch.font.size = Pt(14)
        r_ch.font.bold = True
        r_ch.font.color.rgb = primary_color

        for sec_title, sec_text in chap["sections"]:
            p_sec = doc.add_paragraph()
            p_sec.paragraph_format.space_before = Pt(8)
            p_sec.paragraph_format.space_after = Pt(4)
            r_st = p_sec.add_run(sec_title)
            r_st.font.name = "Arial"
            r_st.font.size = Pt(12)
            r_st.font.bold = True
            r_st.font.color.rgb = dark_color

            p_body = doc.add_paragraph()
            p_body.paragraph_format.line_spacing = 1.15
            p_body.paragraph_format.space_after = Pt(8)
            r_bt = p_body.add_run(sec_text)
            r_bt.font.name = "Arial"
            r_bt.font.size = Pt(10.5)

        # Check for table
        if "table" in chap:
            tbl_info = chap["table"]
            tbl = doc.add_table(rows=len(tbl_info["rows"]) + 1, cols=len(tbl_info["headers"]))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Header Row
            hdr_cells = tbl.rows[0].cells
            for col_idx, h_text in enumerate(tbl_info["headers"]):
                hdr_cells[col_idx].text = h_text
                set_cell_background(hdr_cells[col_idx], "4F46E5")
                set_cell_margins(hdr_cells[col_idx], top=120, bottom=120, left=150, right=150)
                p = hdr_cells[col_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)

            # Data Rows
            for r_idx, row_data in enumerate(tbl_info["rows"]):
                row_cells = tbl.rows[r_idx + 1].cells
                bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx, val in enumerate(row_data):
                    row_cells[c_idx].text = str(val)
                    set_cell_background(row_cells[c_idx], bg_color)
                    set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
                    p = row_cells[c_idx].paragraphs[0]
                    for r in p.runs:
                        r.font.name = "Arial"
                        r.font.size = Pt(9)
                        if val == "PASS":
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(16, 185, 129)

            doc.add_paragraph().paragraph_format.space_after = Pt(12)

        doc.add_page_break()

    # Save Word Document
    output_path = r"D:\1.MCA-PROJECTS\StudentAttendanceApp_Report.docx"
    doc.save(output_path)
    print(f"Report generated successfully at: {output_path}")

if __name__ == "__main__":
    create_report()
